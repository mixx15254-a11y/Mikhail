#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta

# ---------- Данные из фото ----------
SEASONS = [
    {
        "season": "2012-2013",
        "format": "7×7",
        "leagues": [
            {
                "name": "2012-2013 • 7×7",
                "format": "7×7",
                "double": True,
                "start": "2012-09-01",
                "time": "18:00",
                "interval": 7,
                "stadium": "Главная арена",
                "teams": [
                    ("Школа 68", "Шкреба"),
                    ("Крылья-школа №7", "Уразаков"),
                    ("Смена-Кошелев", "Майоров"),
                    ("Школа 64/92", "Тетерин"),
                    ("Школа 139", "Гурин"),
                    ("Профики", "Беляев"),
                    ("Смена-Ведиси", "Куликов"),
                    ("Комета", "Женя"),
                    ("Эверест", "Пузырев"),
                    ("Молот", "Цыганов"),
                    ("Союз", "Цыганов"),
                    ("Школа 1", "Галимуллин"),
                    ("Галатасарай", "Лысиков"),
                    ("Коршуны", "Анисимов"),
                    ("Коршуны 2", "Анисимов"),
                    ("Коршуны 3", "Анисимов"),
                ]
            }
        ]
    },
    {
        "season": "2014-2015",
        "format": "7×7",
        "leagues": [
            {
                "name": "2014-2015 ПРЕМЬЕР-ЛИГА 7×7",
                "format": "7×7",
                "double": True,
                "start": "2014-09-06",
                "time": "18:00",
                "interval": 7,
                "stadium": "Главная арена",
                "teams": [
                    ("Крылья-школа №7", "Уразаков"),
                    ("Смена-Кошелев", "Майоров"),
                    ("Эверест 1", "Пузырев"),
                    ("Профики", "Беляев"),
                    ("КС 26", "Кузьмичев"),
                    ("Школа 1", "Галимуллин"),
                    ("Молния", "Касьянов"),
                    ("Орлы", "Анисимов"),
                ]
            },
            {
                "name": "2014-2015 ВЫСШАЯ ЛИГА 7×7",
                "format": "7×7",
                "double": False,
                "start": "2014-09-07",
                "time": "17:00",
                "interval": 7,
                "stadium": "Поле №2",
                "teams": [
                    ("68 школа", "Шкреба"),
                    ("Школа 175/96", "Лысиков"),
                    ("Крылья дубль-школа №7", "Уразаков"),
                    ("Школа 64/92", "Тетерин"),
                    ("Смена КС (К)", "Якимов"),
                    ("Школа 139", "Гурин"),
                    ("Школа 16", "Смолкин"),
                    ("КС 26", "Кузьмичев"),
                    ("Смена-Ведиси", "Куликов"),
                    ("Комета", "Женя"),
                    ("Молот", "Цыганов"),
                    ("Олимп", "Седин"),
                    ("Школа 35", "Колюжный"),
                    ("Галатасарай", "Лысиков"),
                    ("Орлы 2", "Анисимов"),
                    ("Орлы 3", "Анисимов"),
                    ("Орлы 4", "Анисимов"),
                    ("Орлы 5", "Анисимов"),
                ]
            },
        ]
    },
    {
        "season": "2016-2017",
        "format": "7×7 / 5×5",
        "leagues": [
            {
                "name": "2016-2017 ПРЕМЬЕР-ЛИГА 7×7",
                "format": "7×7",
                "double": True,
                "start": "2016-09-03",
                "time": "18:00",
                "interval": 7,
                "stadium": "Главная арена",
                "teams": [
                    ("68 школа", "Шкреба"),
                    ("Школа 175/96", "Лысиков"),
                    ("Крылья-Школа №7", "Уразаков"),
                    ("Смена-Кошелев", "Майоров"),
                    ("Профики", "Беляев"),
                    ("школа 16", "Смолкин"),
                    ("КС 26", "Кузьмичев"),
                    ("Комета", "Женя"),
                    ("Школа 35", "Колюжный"),
                    ("Красная фурия", "Касьянов"),
                    ("Огонь", "Анисимов"),
                    ("ФШМ", "Винокуров"),
                ]
            },
            {
                "name": "2016-2017 ВЫСШАЯ ЛИГА 5×5 — 5 ГРУПП × 4",
                "format": "5×5",
                "double": False,
                "start": "2016-09-04",
                "time": "16:00",
                "interval": 7,
                "stadium": "Манеж 5×5",
                "groups": 5,
                "group_names": ["Группа A", "Группа B", "Группа C", "Группа D", "Группа E"],
                "teams": [
                    ("Эверест 2", "Пузырев"),
                    ("Школа 64/92", "Тетерин"),
                    ("Школа 139", "Гурин"),
                    ("Профики", "Беляев"),
                    ("Профики-2", "Беляев"),
                    ("школа 16", "Смолкин"),
                    ("КС 26", "Кузьмичев"),
                    ("Комета", "Женя"),
                    ("Молот", "Цыганов"),
                    ("Школа 1", "Галимуллин"),
                    ("Школа 1/2", "Владимиров"),
                    ("Барсы", "Седин"),
                    ("Красная фурия 2", "Касьянов"),
                    ("Олимп 2", "Щадин"),
                    ("Огонь 2", "Анисимов"),
                    ("Огонь 3", "Анисимов"),
                    ("Огонь 4", "Анисимов"),
                    ("Огонь 5", "Анисимов"),
                    ("Дружба", "Майоров"),
                    # добавляем 20-ю чтобы получилось ровно 5×4 — можешь заменить на свою команду
                    ("Эверест 3 (резерв)", "Пузырев"),
                ]
            },
        ]
    },
    {
        "season": "2018-2019",
        "format": "5×5",
        "leagues": [
            {
                "name": "2018-2019 ПРЕМЬЕР-ЛИГА 5×5",
                "format": "5×5",
                "double": True,
                "start": "2018-09-08",
                "time": "18:00",
                "interval": 7,
                "stadium": "Манеж 5×5",
                "teams": [
                    ("68 школа", "Шкреба"),
                    ("Спутник-1", "Агарков"),
                    ("Смена-Кошелев", "Майоров"),
                    ("Комета", "Женя"),
                    ("Молот", "Цыганов"),
                    ("Акулы 42/11", "Щадин"),
                    ("Красные опасные", "Анисимов"),
                    ("Профики", "Беляев"),
                ]
            },
            {
                "name": "2018-2019 ВЫСШАЯ ЛИГА 5×5",
                "format": "5×5",
                "double": False,
                "start": "2018-09-09",
                "time": "17:00",
                "interval": 7,
                "stadium": "Поле 5×5 №2",
                "teams": [
                    ("Акулы 42/11", "Щадин"),
                    ("Школа 64/92", "Тетерин"),
                    ("Профики", "Беляев"),
                    ("Школа 16", "Смолкин"),
                    ("КС 26", "Кузьмичев"),
                    ("Школа 1", "Галимуллин"),
                    ("Школа 35", "Колюжный"),
                    ("Дружба", "Майоров"),
                    ("Красные опасные 2", "Анисимов"),
                    ("Красные опасные 3", "Анисимов"),
                ]
            },
        ]
    },
    {
        "season": "2020-2021",
        "format": "5×5",
        "leagues": [
            {
                "name": "2020-2021 • 13 команд",
                "format": "5×5",
                "double": False,
                "start": "2020-09-05",
                "time": "17:00",
                "interval": 7,
                "stadium": "Главная арена",
                "teams": [
                    ("68 школа", "Шкреба"),
                    ("Спутник", "Агарков"),
                    ("Киты 42/11", "Щадин"),
                    ("Профики", "Беляев"),
                    ("школа 76", "Глинин"),
                    ("Молот (153/157 школа)", "Вукалов"),
                    ("Школа 1", "Галимуллин"),
                    ("Барселона", "Ботов"),
                    ("Крепыши", "Анисимов"),
                    ("Крепыши 2", "Анисимов"),
                    ("Крепыши 3", "Анисимов"),
                    ("ЯнгФорс", "Сусляев"),
                    ("Профики 2", "Беляев"),
                ]
            }
        ]
    },
]

# --- helpers ---
def round_robin(team_names, double=False):
    teams = list(team_names)
    is_odd = len(teams) % 2 == 1
    if is_odd:
        teams.append("— выходной —")
    n = len(teams)
    rounds = n - 1
    half = n // 2
    arr = list(teams)
    schedule = []
    gid = 1
    for r in range(rounds):
        for i in range(half):
            home = arr[i]
            away = arr[n-1-i]
            if home == "— выходной —" or away == "— выходной —":
                continue
            # fairness swap
            is_swap = (r % 2 == 1) and (i == 0)
            if is_swap:
                home, away = away, home
            schedule.append({"id": gid, "round": r+1, "home": home, "away": away})
            gid += 1
        # rotate
        fixed = arr[0]
        rest = arr[1:]
        rest = [rest[-1]] + rest[:-1]
        arr = [fixed] + rest
    if double:
        second = []
        for m in schedule:
            second.append({"id": gid, "round": m["round"] + rounds, "home": m["away"], "away": m["home"]})
            gid += 1
        schedule = schedule + second
    return schedule

# Styles
HEADER_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
HEADER_FONT = Font(name="Calibri", color="FFFFFF", bold=True, size=10)
SUB_HEADER_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
SUB_HEADER_FONT = Font(name="Calibri", color="334155", bold=True, size=9)
TITLE_FILL = PatternFill(start_color="16A34A", end_color="16A34A", fill_type="solid")
TITLE_FONT = Font(name="Calibri", color="FFFFFF", bold=True, size=13)
SECTION_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
EMERALD_FILL = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
thin = Side(style="thin", color="CBD5E1")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header_row(ws, row, cols, fill=HEADER_FILL, font=HEADER_FONT):
    for c in range(1, cols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def auto_width(ws, min_width=10, max_width=28):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    l = len(str(cell.value))
                    if l > max_len:
                        max_len = l
            except:
                pass
        adjusted = min(max_len + 4, max_width)
        if adjusted < min_width:
            adjusted = min_width
        ws.column_dimensions[col_letter].width = adjusted

def create_summary_sheet(wb):
    ws = wb.active
    ws.title = "ОГЛАВЛЕНИЕ"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "1:1"

    # Title
    ws.merge_cells("A1:E1")
    c = ws["A1"]
    c.value = "МИХАИЛ • РАСПИСАНИЕ МАТЧЕЙ • 5 СЕЗОНОВ"
    c.font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:E2")
    c = ws["A2"]
    c.value = "Архивные сезоны 2012 — 2021  •  Форматы 7×7 и 5×5  •  Сгенерировано автоматически"
    c.font = Font(name="Calibri", size=9, color="64748B", italic=True)
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 16

    # Table header
    headers = ["Сезон", "Лига", "Формат", "Команд", "Всего матчей"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[4].height = 22

    row = 5
    flat = []
    for s in SEASONS:
        for lg in s["leagues"]:
            if "groups" in lg:
                groups = split_groups(lg["teams"], lg["groups"])
                total = sum(len(round_robin([t[0] for t in g], lg["double"])) for g in groups)
                flat.append((s["season"], lg["name"], lg["format"], len(lg["teams"]), total))
            else:
                names = [t[0] for t in lg["teams"]]
                sched = round_robin(names, lg["double"])
                flat.append((s["season"], lg["name"], lg["format"], len(lg["teams"]), len(sched)))

    for season, name, fmt, teams_cnt, matches_cnt in flat:
        ws.cell(row=row, column=1, value=season).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=2, value=name).alignment = Alignment(horizontal="left", vertical="center")
        ws.cell(row=row, column=3, value=fmt).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=4, value=teams_cnt).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=5, value=matches_cnt).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=5).font = Font(bold=True, color="16A34A")
        # styles
        is_alt = row % 2 == 0
        for c in range(1,6):
            cell = ws.cell(row=row, column=c)
            cell.border = BORDER
            cell.font = Font(name="Calibri", size=10, color="0F172A", bold=(c==5))
            if is_alt:
                cell.fill = ALT_FILL
            if c==2:
                cell.font = Font(name="Calibri", size=10, bold=True)
            if c==3:
                # format badge color
                if "7×7" in str(cell.value):
                    cell.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                    cell.font = Font(name="Calibri", bold=True, color="0C4A6E", size=10)
                else:
                    cell.fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
                    cell.font = Font(name="Calibri", bold=True, color="92400E", size=10)
        ws.row_dimensions[row].height = 18
        row += 1

    # Notes
    r = row + 2
    ws.merge_cells(f"A{r}:E{r}")
    c = ws.cell(row=r, column=1, value="Как пользоваться:")
    c.font = Font(name="Calibri", bold=True, size=11, color="0F172A")
    notes = [
        "• Каждая лига — отдельный лист (вкладка внизу). Листай вкладки.",
        "• В листе: сверху — список команд и тренеров, ниже — расписание по турам (каждый с каждым).",
        "• Расписание сгенерировано по круговой системе, без повторов, с чередованием дома/в гостях.",
        "• Даты туров — каждую неделю от старта сезона, время и стадион можно менять прямо в Excel.",
        "• Столбцы «Взнос» и «Заявка» оставь для отметок (галочки). Столбец «Счёт» заполняй после игр — таблица считается сама в сайте.",
        "• Для печати: Файл → Печать → Альбомная ориентация, вписать на 1 страницу по ширине.",
    ]
    for i, note in enumerate(notes, 1):
        ws.merge_cells(f"A{r+i}:E{r+i}")
        c = ws.cell(row=r+i, column=1, value=note)
        c.font = Font(name="Calibri", size=9, color="475569")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[r+i].height = 14

    auto_width(ws)
    ws.column_dimensions["B"].width = 34
    ws.freeze_panes = "A5"
    # print
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

def split_groups(teams, n_groups):
    # режем последовательно по 4: Группа A = первые 4, B = следующие 4 и т.д.
    groups = []
    chunk = len(teams) // n_groups
    rem = len(teams) % n_groups
    idx = 0
    for i in range(n_groups):
        size = chunk + (1 if i < rem else 0)
        groups.append(teams[idx:idx+size])
        idx += size
    return groups

def create_league_sheet(wb, league, season_name):
    name = league["name"]
    # sanitize sheet name (max 31 chars)
    sheet_name = name[:31]
    # handle duplicates
    base = sheet_name
    idx = 2
    while sheet_name in wb.sheetnames:
        suffix = f" ({idx})"
        sheet_name = (base[:31-len(suffix)] + suffix)
        idx += 1
    ws = wb.create_sheet(title=sheet_name)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.4
    ws.page_margins.right = 0.4
    ws.page_margins.top = 0.4
    ws.page_margins.bottom = 0.4
    ws.print_title_rows = "1:10"

    is_grouped = "groups" in league
    # Title block
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = name.upper()
    c.font = TITLE_FONT
    c.fill = TITLE_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells("A2:H2")
    c = ws["A2"]
    double_txt = "ДВОЙНОЙ КРУГ (дома/в гостях)" if league["double"] else "ОДИН КРУГ"
    if is_grouped:
        c.value = f"Сезон {season_name}  •  Формат {league['format']}  •  {double_txt}  •  5 ГРУПП × 4 КОМАНДЫ  •  Старт: {league['start']}  •  {league['time']}  •  {league['stadium']}  •  {len(league['teams'])} команд"
    else:
        c.value = f"Сезон {season_name}  •  Формат {league['format']}  •  {double_txt}  •  Старт: {league['start']}  •  {league['time']}  •  {league['stadium']}  •  {len(league['teams'])} команд"
    c.font = Font(name="Calibri", size=8, color="FFFFFF", bold=True)
    c.fill = PatternFill(start_color="15803D", end_color="15803D", fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    ws.merge_cells("A3:H3")
    c = ws["A3"]
    if is_grouped:
        # для групп: каждая группа 4 команды = 3 тура, 6 матчей (одинарный)
        total_matches = 0
        groups = split_groups(league["teams"], league["groups"])
        for g in groups:
            total_matches += len(round_robin([t[0] for t in g], league["double"]))
        c.value = f"Групповой этап: 5 групп по 4 команды  •  В каждой группе каждый с каждым (3 тура, 6 матчей)  •  Всего матчей: {total_matches}  •  Далее плей-офф по желанию"
    else:
        rounds_est = len(league["teams"])-1 if not league["double"] else (len(league["teams"])-1)*2
        if len(league["teams"]) %2==1:
            rounds_est = len(league["teams"]) if not league["double"] else len(league["teams"])*2
        total_matches = len(round_robin([t[0] for t in league["teams"]], league["double"]))
        c.value = f"Всего туров: {rounds_est}  •  Всего матчей: {total_matches}  •  Формула: каждый с каждым  •  Сгенерировано автоматически"
    c.font = Font(name="Calibri", size=8, color="475569", italic=True)
    c.alignment = Alignment(horizontal="center")
    ws.row_dimensions[3].height = 14

    # Teams header
    row = 5
    ws.merge_cells(f"A{row}:H{row}")
    c = ws.cell(row=row, column=1, value="1. КОМАНДЫ И ТРЕНЕРЫ")
    c.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    c.fill = HEADER_FILL
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[row].height = 20

    row += 1
    headers = ["№", "Команда", "Тренер", "Взнос", "Заявка", "", "", ""]
    # Actually 5 columns: №, Команда, Тренер, Взнос, Заявка -> we use A-E, but need nice widths
    team_headers = ["№", "Команда", "Тренер", "Взнос", "Заявка"]
    for col, h in enumerate(team_headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = SUB_HEADER_FILL
        cell.font = SUB_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.row_dimensions[row].height = 18

    # Teams rows
    row += 1
    start_team_row = row
    for idx, (team, trainer) in enumerate(league["teams"], 1):
        ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row, column=1).font = Font(name="Calibri", bold=True, color="0F172A")
        ws.cell(row=row, column=1).border = BORDER
        # team
        c = ws.cell(row=row, column=2, value=team)
        c.font = Font(name="Calibri", bold=True, size=10)
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BORDER
        # trainer
        c = ws.cell(row=row, column=3, value=trainer)
        c.font = Font(name="Calibri", size=10, color="475569")
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BORDER
        # vznos
        c = ws.cell(row=row, column=4, value="☐")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(size=12)
        c.border = BORDER
        # zayavka
        c = ws.cell(row=row, column=5, value="☐")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.font = Font(size=12)
        c.border = BORDER
        if idx % 2 == 0:
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = ALT_FILL
        else:
            for col in range(1, 6):
                ws.cell(row=row, column=col).fill = PatternFill(fill_type=None)
        ws.row_dimensions[row].height = 16
        row += 1
    end_team_row = row - 1

    # --------- РАСПИСАНИЕ ---------
    if is_grouped:
        groups = split_groups(league["teams"], league["groups"])
        group_names = league.get("group_names", [f"Группа {chr(65+i)}" for i in range(len(groups))])
        base_date = datetime.strptime(league["start"] + " " + league["time"], "%Y-%m-%d %H:%M")
        interval = league["interval"]
        header_row = None
        row += 1
        ws.merge_cells(f"A{row}:H{row}")
        c = ws.cell(row=row, column=1, value="2. ГРУППОВОЙ ЭТАП — 5 ГРУПП ПО 4 КОМАНДЫ (каждая группа — круговая)")
        c.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 20
        for gi, group in enumerate(groups):
            gname = group_names[gi] if gi < len(group_names) else f"Группа {chr(65+gi)}"
            # Group title
            row += 1
            ws.merge_cells(f"A{row}:H{row}")
            c = ws.cell(row=row, column=1, value=f"{gname.upper()} — {len(group)} команды: {', '.join([t[0] for t in group])}")
            c.font = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
            # цвет группы
            colors = ["0F172A","1D4ED8","065F46","7C2D12","6D28D9"]
            col = colors[gi % len(colors)]
            c.fill = PatternFill(start_color=col, end_color=col, fill_type="solid")
            c.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row].height = 18
            # Teams of group mini-table
            row += 1
            for col_idx, h in enumerate(["№","Команда","Тренер","Взнос","Заявка"], 1):
                cell = ws.cell(row=row, column=col_idx, value=h)
                cell.fill = SUB_HEADER_FILL
                cell.font = SUB_HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = BORDER
            ws.row_dimensions[row].height = 16
            for idx, (team, trainer) in enumerate(group, 1):
                row += 1
                ws.cell(row=row, column=1, value=idx).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=1).font = Font(bold=True, size=9)
                ws.cell(row=row, column=1).border = BORDER
                c = ws.cell(row=row, column=2, value=team)
                c.font = Font(bold=True, size=9)
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.border = BORDER
                c = ws.cell(row=row, column=3, value=trainer)
                c.font = Font(size=9, color="475569")
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.border = BORDER
                ws.cell(row=row, column=4, value="☐").alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=4).border = BORDER
                ws.cell(row=row, column=5, value="☐").alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=5).border = BORDER
                if idx % 2 == 0:
                    for cc in range(1,6):
                        ws.cell(row=row, column=cc).fill = ALT_FILL
                ws.row_dimensions[row].height = 14
            # Schedule for this group
            row += 1
            ws.merge_cells(f"A{row}:H{row}")
            c = ws.cell(row=row, column=1, value=f"Расписание {gname} — каждый с каждым")
            c.font = Font(name="Calibri", bold=True, size=9, color="334155")
            c.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = BORDER
            ws.row_dimensions[row].height = 16
            row += 1
            for col_idx, h in enumerate(["Тур","Дата","Время","Хозяева","—","Гости","Стадион","Счёт"], 1):
                cell = ws.cell(row=row, column=col_idx, value=h)
                cell.fill = SUB_HEADER_FILL
                cell.font = SUB_HEADER_FONT
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = BORDER
            cur_header = row
            if header_row is None:
                header_row = cur_header
            ws.row_dimensions[row].height = 16
            # generate group schedule
            g_names = [t[0] for t in group]
            sched = round_robin(g_names, league["double"])
            for m in sched:
                d = base_date + timedelta(days=(m["round"]-1)*interval)
                m["date"] = d.strftime("%d.%m.%Y")
                m["time_val"] = league["time"]
                m["stadium"] = league["stadium"] + f" ({gname})"
            # для каждой группы туры отдельно (1-3)
            for m in sorted(sched, key=lambda x: (x["round"], x["id"])):
                row += 1
                ws.cell(row=row, column=1, value=m["round"]).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=1).font = Font(bold=True, size=9)
                ws.cell(row=row, column=1).border = BORDER
                ws.cell(row=row, column=2, value=m["date"]).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=2).font = Font(size=9)
                ws.cell(row=row, column=2).border = BORDER
                ws.cell(row=row, column=3, value=m["time_val"]).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=3).font = Font(size=9)
                ws.cell(row=row, column=3).border = BORDER
                c = ws.cell(row=row, column=4, value=m["home"])
                c.font = Font(bold=True, size=9)
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.border = BORDER
                ws.cell(row=row, column=5, value="—").alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=5).border = BORDER
                c = ws.cell(row=row, column=6, value=m["away"])
                c.font = Font(bold=True, size=9)
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.border = BORDER
                ws.cell(row=row, column=7, value=m["stadium"]).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=7).font = Font(size=7, color="475569")
                ws.cell(row=row, column=7).border = BORDER
                ws.cell(row=row, column=8, value=" : ").alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=row, column=8).font = Font(bold=True, size=10)
                ws.cell(row=row, column=8).border = BORDER
                if row % 2 == 0:
                    for cc in range(1,9):
                        if ws.cell(row=row, column=cc).fill == PatternFill(fill_type=None):
                            ws.cell(row=row, column=cc).fill = ALT_FILL
                ws.row_dimensions[row].height = 14
            row += 1  # gap between groups
        # Footer
        row += 1
        ws.merge_cells(f"A{row}:H{row}")
        c = ws.cell(row=row, column=1, value="* В каждой группе — 3 тура, 6 матчей (один круг). Для выхода в плей-офф бери по 1-2 лучшие команды из каждой группы. Счёт вписывай в столбец «Счёт». Группу E можешь переименовать, резерв убери если не нужен.")
        c.font = Font(name="Calibri", size=8, italic=True, color="64748B")
        c.alignment = Alignment(horizontal="left", wrap_text=True)
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 9
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 4
        ws.column_dimensions["F"].width = 22
        ws.column_dimensions["G"].width = 18
        ws.column_dimensions["H"].width = 9
        ws.freeze_panes = "A8"
        ws.print_area = f"A1:H{row}"
        if header_row:
            ws.auto_filter.ref = f"A{header_row}:H{row-2}"
        return ws
    else:
        # обычный формат без групп
        names = [t[0] for t in league["teams"]]
        schedule = round_robin(names, league["double"])
        base_date = datetime.strptime(league["start"] + " " + league["time"], "%Y-%m-%d %H:%M")
        interval = league["interval"]
        for m in schedule:
            d = base_date + timedelta(days=(m["round"]-1)*interval)
            m["date"] = d.strftime("%d.%m.%Y")
            m["time_val"] = league["time"]
            m["stadium"] = league["stadium"]
            m["status"] = ""
            m["score"] = ""
        row += 1
        ws.merge_cells(f"A{row}:H{row}")
        c = ws.cell(row=row, column=1, value="2. РАСПИСАНИЕ МАТЧЕЙ — КРУГОВАЯ СИСТЕМА (каждый с каждым)")
        c.font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row].height = 20
        row += 1
        sched_headers = ["Тур", "Дата", "Время", "Хозяева", "", "Гости", "Стадион", "Счёт"]
        for col, h in enumerate(sched_headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.fill = SUB_HEADER_FILL if h != "" else SUB_HEADER_FILL
            cell.font = SUB_HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = BORDER
            if col == 5:
                cell.value = "—"
                cell.font = Font(name="Calibri", size=9, color="94A3B8")
        ws.row_dimensions[row].height = 20
        header_row = row
        row += 1
        current_round = None
        for m in sorted(schedule, key=lambda x: (x["round"], x["id"])):
            if m["round"] != current_round:
                ws.merge_cells(f"A{row}:H{row}")
                c = ws.cell(row=row, column=1, value=f"ТУР {m['round']}  —  {m['date']}")
                c.font = Font(name="Calibri", bold=True, size=9, color="0F172A")
                c.fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.border = BORDER
                for col in range(1, 9):
                    ws.cell(row=row, column=col).fill = PatternFill(start_color="E0F2FE", end_color="E0F2FE", fill_type="solid")
                    ws.cell(row=row, column=col).border = BORDER
                ws.row_dimensions[row].height = 18
                row += 1
                current_round = m["round"]
            ws.cell(row=row, column=1, value=m["round"]).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=1).font = Font(name="Calibri", bold=True, size=9)
            ws.cell(row=row, column=1).border = BORDER
            ws.cell(row=row, column=2, value=m["date"]).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=2).font = Font(name="Calibri", size=9)
            ws.cell(row=row, column=2).border = BORDER
            ws.cell(row=row, column=3, value=m["time_val"]).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=3).font = Font(name="Calibri", size=9)
            ws.cell(row=row, column=3).border = BORDER
            c = ws.cell(row=row, column=4, value=m["home"])
            c.font = Font(name="Calibri", bold=True, size=10)
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.border = BORDER
            ws.cell(row=row, column=5, value="—").alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=5).font = Font(color="94A3B8", size=9)
            ws.cell(row=row, column=5).border = BORDER
            c = ws.cell(row=row, column=6, value=m["away"])
            c.font = Font(name="Calibri", bold=True, size=10)
            c.alignment = Alignment(horizontal="left", vertical="center")
            c.border = BORDER
            ws.cell(row=row, column=7, value=m["stadium"]).alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=7).font = Font(size=8, color="475569")
            ws.cell(row=row, column=7).border = BORDER
            ws.cell(row=row, column=8, value=" : ").alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(row=row, column=8).font = Font(bold=True, size=11)
            ws.cell(row=row, column=8).border = BORDER
            if row % 2 == 0:
                for col in range(1, 9):
                    if ws.cell(row=row, column=col).fill == PatternFill(fill_type=None):
                        ws.cell(row=row, column=col).fill = ALT_FILL
            ws.row_dimensions[row].height = 15
            row += 1
        row += 1
        ws.merge_cells(f"A{row}:H{row}")
        c = ws.cell(row=row, column=1, value="* Счёт вписывай в столбец «Счёт» (напр. 2:1). Для переноса матча меняй дату/время. Распечатай: Файл → Печать → Альбомная, вписать на 1 стр.")
        c.font = Font(name="Calibri", size=8, italic=True, color="64748B")
        c.alignment = Alignment(horizontal="left", wrap_text=True)
        ws.column_dimensions["A"].width = 6
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 9
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 4
        ws.column_dimensions["F"].width = 22
        ws.column_dimensions["G"].width = 16
        ws.column_dimensions["H"].width = 9
        ws.freeze_panes = "A8"
        ws.print_area = f"A1:H{row}"
        ws.auto_filter.ref = f"A{header_row}:H{row-2}"
        return ws

def main():
    wb = openpyxl.Workbook()
    create_summary_sheet(wb)

    for s in SEASONS:
        for lg in s["leagues"]:
            create_league_sheet(wb, lg, s["season"])

    # Set active to summary
    wb.active = 0

    out_main = "Mikhail_5_Sezonov_Raspisanie.xlsx"
    wb.save(out_main)
    print(f"Saved {out_main} with {len(wb.sheetnames)} sheets: {wb.sheetnames}")

    # Also create separate files per season for convenience?
    for s in SEASONS:
        wb2 = openpyxl.Workbook()
        # remove default
        wb2.remove(wb2.active)
        # summary for this season only
        # create a mini summary?
        for lg in s["leagues"]:
            create_league_sheet(wb2, lg, s["season"])
        fname = f"Mikhail_{s['season'].replace('-','_')}_{s['leagues'][0]['format'].replace('×','x')}.xlsx"
        # sanitize
        fname = fname.replace("/", "_").replace(" ", "_")
        wb2.save(fname)
        print(f"Saved {fname}")

if __name__ == "__main__":
    main()
