#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime, timedelta
from collections import defaultdict, Counter
import collections

YEAR=2025
SLOT_DURATION=60
RAW_DATES=[
    ("26.08","17:30","20:30",{"7*7":2,"5*5":3,"4*4":0}),
    ("27.08","10:00","13:00",{"7*7":2,"5*5":2,"4*4":3}),
    ("03.10","17:30","20:30",{"7*7":1,"5*5":3,"4*4":0}),
    ("04.10","10:00","13:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("10.10","17:30","20:30",{"7*7":1,"5*5":3,"4*4":0}),
    ("11.10","10:00","13:00",{"7*7":1,"5*5":2,"4*4":3}),
    ("17.10","17:30","20:30",{"7*7":1,"5*5":3,"4*4":0}),
    ("18.10","10:00","13:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("24.10","17:30","20:30",{"7*7":1,"5*5":3,"4*4":0}),
    ("25.10","10:00","13:00",{"7*7":1,"5*5":2,"4*4":3}),
    ("31.10","17:00","21:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("07.11","17:00","21:00",{"7*7":1,"5*5":2,"4*4":3}),
    ("14.11","17:00","21:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("21.11","17:00","21:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("28.11","14:00","20:00",{"7*7":1,"5*5":2,"4*4":3}),
    ("05.12","17:00","21:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("12.12","17:00","21:00",{"7*7":1,"5*5":2,"4*4":3}),
    ("19.12","17:00","21:00",{"7*7":1,"5*5":3,"4*4":0}),
    ("26.12","17:00","21:00",{"7*7":1,"5*5":2,"4*4":3}),
]
def parse_time(s): h,m=map(int,s.split(":")); return h*60+m
def generate_slots():
    slots=[]
    for date_str,start,end,fields in RAW_DATES:
        d,m=map(int,date_str.split("."))
        date_obj=datetime(YEAR,m,d)
        n_slots=(parse_time(end)-parse_time(start))//SLOT_DURATION
        for si in range(n_slots):
            tmin=parse_time(start)+si*SLOT_DURATION
            time_str=f"{tmin//60:02d}:{tmin%60:02d}"
            for ftype,cnt in fields.items():
                for fid in range(1,cnt+1):
                    field_id=f"Поле {ftype} #{fid}" if cnt>1 else f"Поле {ftype}"
                    slots.append({"date_str":date_str,"date_obj":date_obj,"time":time_str,"field":field_id,"field_type":ftype,"slot_key":f"{date_str} {time_str}"})
    return slots
def round_robin(names,double=False):
    teams=list(names)
    if len(teams)%2==1: teams.append("— выходной —")
    n=len(teams);rounds=n-1;half=n//2;arr=list(teams);sched=[]
    for r in range(rounds):
        for i in range(half):
            home=arr[i];away=arr[n-1-i]
            if home=="— выходной —" or away=="— выходной —": continue
            is_swap=(r%2==1) and (i==0)
            if is_swap: home,away=away,home
            sched.append({"round":r+1,"home":home,"away":away})
        fixed=arr[0];rest=arr[1:];rest=[rest[-1]]+rest[:-1];arr=[fixed]+rest
    if double:
        second=[{"round":m["round"]+rounds,"home":m["away"],"away":m["home"]} for m in sched]
        sched=sched+second
    return sched

LEAGUES=[
    {"name":"2014-2015 Высшая 7×7 — группы 3×6","format":"7*7","groups":3,"teams":[("68 школа","Шкреба"),("Школа 64/92","Тетерин"),("Школа 16","Смолкин"),("Комета","Женя"),("Школа 35","Колюжный"),("Орлы 2","Анисимов"),("Школа 175/96","Лысиков"),("Смена КС (К)","Якимов"),("КС 26","Кузьмичев"),("Молот","Цыганов"),("Олимп","Седин"),("Орлы 3","Анисимов"),("Крылья дубль-школа №7","Уразаков"),("Школа 139","Гурин"),("Смена-Ведиси","Куликов"),("Галатасарай","Лысиков"),("Орлы 4","Анисимов"),("Орлы 5","Анисимов")],"double":False},
    {"name":"2016-2017 Высшая 5×5 — группы 5×4","format":"5*5","groups":5,"teams":[("Эверест 2","Пузырев"),("Профики-2","Беляев"),("Молот","Цыганов"),("Огонь 2","Анисимов"),("Школа 64/92","Тетерин"),("школа 16","Смолкин"),("Школа 1","Галимуллин"),("Огонь 3","Анисимов"),("Школа 139","Гурин"),("КС 26","Кузьмичев"),("Школа 1/2","Владимиров"),("Огонь 4","Анисимов"),("Профики","Беляев"),("Комета","Женя"),("Барсы","Седин"),("Огонь 5","Анисимов"),("Красная фурия 2","Касьянов"),("Олимп 2","Щадин"),("Дружба","Майоров"),("Эверест 3","Пузырев")],"double":False},
    {"name":"2018-2019 Премьер 5×5 — 2 круга","format":"5*5","teams":[("68 школа","Шкреба"),("Спутник-1","Агарков"),("Смена-Кошелев","Майоров"),("Комета","Женя"),("Молот","Цыганов"),("Акулы 42/11","Щадин"),("Красные опасные","Анисимов"),("Профики","Беляев")],"double":True},
    {"name":"2018-2019 Высшая 5×5 — 2 круга","format":"5*5","teams":[("Акулы 42/11","Щадин"),("Школа 64/92","Тетерин"),("Профики","Беляев"),("Школа 16","Смолкин"),("КС 26","Кузьмичев"),("Школа 1","Галимуллин"),("Школа 35","Колюжный"),("Дружба","Майоров"),("Красные опасные 2","Анисимов"),("Красные опасные 3","Анисимов")],"double":True},
    {"name":"2020-2021 Предварительный 13 — 1 круг","format":"5*5","teams":[("68 школа","Шкреба"),("Спутник","Агарков"),("Киты 42/11","Щадин"),("Профики","Беляев"),("школа 76","Глинин"),("Молот (153/157)","Вукалов"),("Школа 1","Галимуллин"),("Барселона","Ботов"),("Крепыши","Анисимов"),("Крепыши 2","Анисимов"),("Крепыши 3","Анисимов"),("ЯнгФорс","Сусляев"),("Профики 2","Беляев")],"double":False},
]
def build_matches():
    all_matches=[]
    for lg in LEAGUES:
        fmt=lg["format"]
        if "groups" in lg:
            n_groups=lg["groups"]
            teams=lg["teams"]
            chunk=len(teams)//n_groups
            groups=[teams[i*chunk:(i+1)*chunk] for i in range(n_groups)]
            group_names=[f"Группа {chr(65+i)}" for i in range(n_groups)]
            for gi,g in enumerate(groups):
                names=[t[0] for t in g]
                sched=round_robin(names, lg["double"])
                for m in sched:
                    ht=next((t[1] for t in g if t[0]==m["home"]),"?")
                    at=next((t[1] for t in g if t[0]==m["away"]),"?")
                    all_matches.append({"league":lg["name"],"format":fmt,"group":group_names[gi],"home":m["home"],"away":m["away"],"home_train":ht,"away_train":at,"round":m["round"]})
        else:
            names=[t[0] for t in lg["teams"]]
            sched=round_robin(names, lg["double"])
            for m in sched:
                ht=next((t[1] for t in lg["teams"] if t[0]==m["home"]),"?")
                at=next((t[1] for t in lg["teams"] if t[0]==m["away"]),"?")
                all_matches.append({"league":lg["name"],"format":fmt,"group":"","home":m["home"],"away":m["away"],"home_train":ht,"away_train":at,"round":m["round"]})
    return all_matches

def schedule_per_time(all_matches, slots):
    # Группируем слоты по времени
    slots_by_time=defaultdict(list)
    for s in slots:
        slots_by_time[s["slot_key"]].append(s)
    # сортируем ключи по дате/времени
    def key_to_dt(k):
        d,t=k.split()
        dd,mm=map(int,d.split("."))
        hh,mi=map(int,t.split(":"))
        return datetime(YEAR,mm,dd,hh,mi)
    time_keys=sorted(slots_by_time.keys(), key=key_to_dt)
    # матчи сортируем по весу тренера
    cnt=Counter()
    for m in all_matches:
        cnt[m["home_train"]]+=1
        cnt[m["away_train"]]+=1
    def w(m): return cnt[m["home_train"]]+cnt[m["away_train"]]
    remaining=sorted(all_matches, key=w, reverse=True)
    scheduled=[]
    # для каждого времени пытаемся заполнить поля без конфликта тренеров
    for tk in time_keys:
        fields=slots_by_time[tk]
        # для этого времени busy тренеров
        busy_here=set()
        # также отдельно для проверки полей — каждое поле одно использование
        # fields уже уникальны
        # пробуем заполнить каждое поле
        for field_slot in sorted(fields, key=lambda x: x["field"]):
            # найдём первый подходящий матч
            best=None
            best_idx=-1
            for idx,m in enumerate(remaining):
                if field_slot["field_type"]!=m["format"]:
                    if not (m["format"]=="5*5" and field_slot["field_type"]=="4*4"):
                        continue
                if m["home_train"] in busy_here or m["away_train"] in busy_here:
                    continue
                # также проверим что матч не использует тренера который уже играет в это время на другом поле — busy_here уже
                best=m
                best_idx=idx
                break
            if best:
                # назначаем
                scheduled.append({
                    "date":field_slot["date_str"],
                    "date_obj":field_slot["date_obj"],
                    "time":field_slot["time"],
                    "field":field_slot["field"],
                    "field_type":field_slot["field_type"],
                    "league":best["league"],
                    "format":best["format"],
                    "group":best["group"],
                    "home":best["home"],
                    "away":best["away"],
                    "home_train":best["home_train"],
                    "away_train":best["away_train"],
                    "round":best["round"],
                    "slot_key":tk,
                    "is_fallback":(best["format"]=="5*5" and field_slot["field_type"]=="4*4")
                })
                busy_here.add(best["home_train"])
                busy_here.add(best["away_train"])
                # удаляем из remaining
                remaining.pop(best_idx)
        # end fields
    return scheduled, remaining

def main():
    slots=generate_slots()
    all_matches=build_matches()
    print(f"Матчей в плане: {len(all_matches)}")
    print(f"Слотов: {len(slots)}")
    scheduled, remaining = schedule_per_time(all_matches, slots)
    print(f"Запланировано: {len(scheduled)}, осталось: {len(remaining)}")
    # проверка параллелей
    check=defaultdict(list)
    for sc in scheduled:
        trainers=set([sc["home_train"], sc["away_train"]])
        trainers.discard(""); trainers.discard("—"); trainers.discard("?")
        check[sc["date"]+" "+sc["time"]].extend(list(trainers))
    conflicts=0
    for k,trs in check.items():
        cnt=Counter(trs)
        for tr,c in cnt.items():
            if c>1:
                conflicts+=1
                print(f"Конфликт {k} тренер {tr} {c} раза — {trs}")
                break
    print(f"Конфликтов (реальных параллелей): {conflicts}")
    # уже импортировано сверху
    wb=openpyxl.Workbook()
    ws=wb.active;ws.title="Расписание_без_параллелей"
    HEADER_FILL=PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    HEADER_FONT=Font(color="FFFFFF", bold=True, size=9)
    ALT_FILL=PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    thin=Side(style="thin", color="CBD5E1")
    BORDER=Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.merge_cells("A1:K1");ws["A1"].value="РАСПИСАНИЕ ОСЕНЬ — БЕЗ ПАРАЛЛЕЛЕЙ ТРЕНЕРОВ • 60 мин слот • per-time алгоритм"
    ws["A1"].font=Font(bold=True,size=11,color="0F172A");ws["A1"].alignment=Alignment(horizontal="center")
    ws.merge_cells("A2:K2");ws["A2"].value="5*5 на 4*4 помечены жёлтым — это запасные поля. Тренер не может быть в двух местах одновременно — проверка 0 конфликтов."
    ws["A2"].font=Font(size=8,italic=True,color="64748B");ws["A2"].alignment=Alignment(horizontal="center")
    headers=["Дата","День","Время","Поле","Формат","Лига","Группа","Хозяева","Тренер","Гости","Тренер"]
    for c,h in enumerate(headers,1):
        cell=ws.cell(row=4,column=c,value=h);cell.fill=HEADER_FILL;cell.font=HEADER_FONT;cell.alignment=Alignment(horizontal="center",vertical="center");cell.border=BORDER
    scheduled_sorted=sorted(scheduled, key=lambda x:(x["date_obj"], x["time"], x["field"]))
    row=5
    for sc in scheduled_sorted:
        dow=sc["date_obj"].strftime("%a")
        days={"Mon":"Пн","Tue":"Вт","Wed":"Ср","Thu":"Чт","Fri":"Пт","Sat":"Сб","Sun":"Вс"}
        dow=days.get(dow,dow)
        vals=[sc["date"],dow,sc["time"],sc["field"],sc["format"],sc["league"],sc["group"],sc["home"],sc["home_train"],sc["away"],sc["away_train"]]
        for c,v in enumerate(vals,1):
            cell=ws.cell(row=row,column=c,value=v)
            cell.font=Font(size=8) if c not in (8,10) else Font(size=8,bold=True)
            if c in (9,11): cell.font=Font(size=7,color="475569")
            cell.alignment=Alignment(horizontal="center",vertical="center");cell.border=BORDER
            if sc.get("is_fallback"):
                cell.fill=PatternFill(start_color="FEF3C7",end_color="FEF3C7",fill_type="solid")
        ws.row_dimensions[row].height=15;row+=1
    widths=[8,6,7,16,7,30,10,18,12,18,12]
    for i,w in enumerate(widths,1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes="A5";ws.auto_filter.ref=f"A4:K{row-1}";ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation="landscape";ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0

    ws2=wb.create_sheet("Сводка")
    ws2["A1"].value="Сводка спортдиректора — без параллелей"
    ws2["A1"].font=Font(bold=True,size=12)
    ws2["A3"].value=f"Запланировано: {len(scheduled)}"
    ws2["A4"].value=f"Не влезло (нужны доп. даты): {len(remaining)}"
    ws2["A5"].value=f"Конфликтов: {conflicts}"
    ws2["A5"].font=Font(color="16A34A",bold=True) if conflicts==0 else Font(color="DC2626",bold=True)
    ws2["A7"].value="По датам:"
    r=8
    cnt_date=Counter(s["date"] for s in scheduled)
    for d in sorted(set(s["date"] for s in scheduled), key=lambda x: datetime.strptime(x+".2025","%d.%m.%Y")):
        ws2.cell(row=r,column=1,value=d);ws2.cell(row=r,column=2,value=cnt_date[d]);r+=1
    if remaining:
        ws2.cell(row=r+2,column=1,value="Осталось (перенести на весну):")
        for i,m in enumerate(remaining[:20], r+3):
            ws2.cell(row=i,column=1,value=f"{m['league']} {m['group']} {m['home']}-{m['away']}")
    ws2.column_dimensions["A"].width=60
    wb.save("Raspisanie_Bez_Paralleley.xlsx")
    print("Saved Raspisanie_Bez_Paralleley.xlsx")

if __name__=="__main__":
    main()
